import streamlit as st
import google.generativeai as genai
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
import io

# 1. إعدادات الصفحة (يجب أن تكون في السطر الأول)
st.set_page_config(page_title="محرك البحث عن الحقيقة", page_icon="⚖️", layout="wide")

# 2. تصميم CSS فاخر (Typography وزخارف بصرية)
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Amiri:wght@400;700&family=Tajawal:wght@400;500;700&display=swap');
    
    /* إعدادات الخلفية والخطوط الأساسية */
    html, body, [class*="css"] {
        font-family: 'Tajawal', sans-serif !important;
        direction: rtl;
        text-align: right;
    }
    
    /* العناوين بخط أميري للمسة كلاسيكية فخمة */
    h1, h2, h3 {
        font-family: 'Amiri', serif !important;
        color: #b8860b !important; /* لون ذهبي داكن */
    }
    
    /* تنسيق الزر الرئيسي */
    .stButton>button {
        background: linear-gradient(135deg, #1e3a8a, #3b82f6);
        color: white;
        border-radius: 8px;
        padding: 10px 24px;
        font-size: 18px;
        font-weight: bold;
        font-family: 'Tajawal', sans-serif;
        border: 1px solid #b8860b;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
        transition: 0.3s;
        width: 100%;
    }
    .stButton>button:hover {
        background: linear-gradient(135deg, #b8860b, #daa520);
        color: white;
        transform: translateY(-2px);
        border-color: #fff;
    }
    
    /* تنسيق صندوق الإدخال */
    .stTextArea textarea {
        background-color: #f8f9fa;
        border: 2px solid #bdc3c7;
        border-radius: 10px;
        font-size: 18px;
    }
    .stTextArea textarea:focus {
        border-color: #b8860b;
        box-shadow: 0 0 5px rgba(184, 134, 11, 0.5);
    }
    </style>
""", unsafe_allow_html=True)

# 3. إعداد الذكاء الاصطناعي
API_KEY = st.secrets["GOOGLE_API_KEY"]
genai.configure(api_key=API_KEY)

system_prompt = """أنت محرك بحث عن الحقيقة متخصص في تحليل الشبهات المتعلقة بالإسلام.
لا تفترض النتيجة مسبقاً، بل اتبع الخطوات الـ 11 بصرامة:
1. تحديد الدعوى. 2. أقوى صياغة للاعتراض. 3. تحليل الافتراضات. 4. بناء خريطة الحجة.
5. جمع الأدلة المتاحة. 6. نقد المصادر. 7. التقييم الأولي. 8. الهجوم. 9. الدفاع. 10. الهجوم المضاد. 11. التقييم النهائي.

ممنوع اختلاق أي مصدر. مسموح لك تماماً بقول "لا نعرف" أو "الأدلة غير كافية" عند الحاجة.
"""
model = genai.GenerativeModel("gemini-3.8-flash")

# دالة إنشاء ملف الوورد في الذاكرة
def generate_word_doc(query, text):
    doc = docx.Document()
    
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(12)

    title = doc.add_heading('تقرير محرك البحث عن الحقيقة 🔎', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_query = doc.add_paragraph()
    p_query.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p_query.add_run(f"الشبهة: {query}")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(192, 57, 43)
    
    doc.add_paragraph('_' * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = text.split('\n')
    for line in lines:
        line = line.strip()
        if not line: continue
        
        if line.startswith('### '):
            p = doc.add_heading(line.replace('### ', ''), level=3)
        elif line.startswith('## '):
            p = doc.add_heading(line.replace('## ', ''), level=2)
        elif line.startswith('# '):
            p = doc.add_heading(line.replace('# ', ''), level=1)
        else:
            p = doc.add_paragraph()
            if line.startswith('- ') or line.startswith('* '):
                line = line.replace('- ', '• ').replace('* ', '• ', 1)
            
            parts = line.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0: run.bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # حفظ الملف في الذاكرة المؤقتة لتحميله مباشرة عبر المتصفح
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# 4. الشريط الجانبي (مكتبة الكتب)
# 4. الشريط الجانبي (مكتبة الكتب)
with st.sidebar:
    st.image("watermarked_img_4216407252451230343.jpg", width=200)
    st.header("📚 مكتبة الردود والشبهات")
    st.write("مكتبة رقمية لتحميل أحدث الكتب والأبحاث.")
    st.markdown("---")
    st.download_button(label="📥 تحميل كتاب: حصون العقيدة (PDF)", data="ملف تجريبي", file_name="book1.pdf", mime="application/pdf")
    st.download_button(label="📥 تحميل كتاب: التفكير الناقد (PDF)", data="ملف تجريبي", file_name="book2.pdf", mime="application/pdf")

# 5. الواجهة الرئيسية
st.title("⚖️ محرك البحث عن الحقيقة")
st.markdown("### 🧠 نظام ذكاء اصطناعي للتحليل المنهجي الدقيق")

st.write("أدخل الشبهة أو السؤال الذي تود تحليله منهجياً في الصندوق أدناه:")
user_input = st.text_area("الشبهة / السؤال:", height=100, placeholder="اكتب الشبهة هنا...")

if st.button("🚀 بدء التحليل المنهجي"):
    if user_input.strip():
        with st.spinner('جاري فحص المصادر وتطبيق منهجية الـ 11 خطوة... لحظات من فضلك ⏳'):
            try:
                full_prompt = f"{system_prompt}\n\n{'-'*50}\nالشبهة المطلوب تحليلها:\n{user_input}"
                response = model.generate_content(full_prompt)
                
                st.success("✅ اكتمل التحليل بنجاح!")
                
                # إنشاء ملف الوورد في الخلفية
                doc_file = generate_word_doc(user_input, response.text)
                
                # زر التحميل بصيغة Word
                st.download_button(
                    label="💾 حفظ التقرير كملف Word احترافي",
                    data=doc_file,
                    file_name="Truth_Report.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
                
                # عرض النتيجة على الشاشة
                st.markdown("---")
                st.markdown(response.text)
                
            except Exception as e:
                st.error(f"حدث خطأ أثناء الاتصال: {e}")
    else:
        st.warning("⚠️ يرجى إدخال الشبهة أولاً قبل الضغط على زر التحليل.")
